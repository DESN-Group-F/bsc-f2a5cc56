"""Safe local readers. They never execute archive code, macros, or MATLAB objects."""
import argparse, csv, json, pathlib, zipfile, tarfile, collections, io, re
class NonfiniteTaggedView:
    NATIVE=b'__READER_NATIVE_STRING_4f7c__:'
    TOK={b'NaN':b'{"__reader_nonfinite__":"NaN"}',b'Infinity':b'{"__reader_nonfinite__":"+Infinity"}',b'-Infinity':b'{"__reader_nonfinite__":"-Infinity"}'}
    def __init__(self,raw): self.raw=raw; self.buf=bytearray(); self.word=bytearray(); self.string=False; self.escape=False; self.eof=False; self.counts=collections.Counter()
    def emit(self):
        if self.word:
            w=bytes(self.word); self.buf.extend(self.TOK.get(w,w))
            if w in self.TOK:self.counts[w.decode('ascii')]+=1
            self.word.clear()
    def fill(self,n):
        while len(self.buf)<n and not self.eof:
            c=self.raw.read(1048576)
            if not c:self.emit();self.eof=True;break
            for b in c:
                if self.string:
                    self.buf.append(b)
                    if self.escape:self.escape=False
                    elif b==92:self.escape=True
                    elif b==34:self.string=False
                elif b==34:self.emit();self.buf.append(b);self.buf.extend(self.NATIVE);self.string=True
                elif 65<=b<=90 or 97<=b<=122 or b==45:self.word.append(b)
                else:self.emit();self.buf.append(b)
    def read(self,n=-1):
        if n<0:
            while not self.eof:self.fill(max(len(self.buf)+1,1048576))
            n=len(self.buf)
        else:self.fill(n)
        r=bytes(self.buf[:n]);del self.buf[:n];return r

def csv_rows(path, limit=1000):
    with open(path,encoding='utf-8-sig',newline='') as f:
        r=csv.DictReader(f); prev=None; segment=0
        for i,row in enumerate(r):
            if i>=limit: break
            cur=None
            try: cur=float(row.get('Test_Time')) if row.get('Test_Time') not in ('',None) else None
            except ValueError: pass
            back=cur is not None and prev is not None and cur<prev
            if back:segment+=1
            out={k:(None if v=='' else v) for k,v in row.items()};out['_reader_segment_id']=segment;out['_reader_test_time_backstep']=back
            yield out
            if cur is not None:prev=cur
def list_members(path):
    p=pathlib.Path(path); s=p.name.lower()
    if zipfile.is_zipfile(p):
        with zipfile.ZipFile(p) as z:
            return [{'path':x.filename,'bytes':x.file_size,'compressed_bytes':x.compress_size,'directory':x.is_dir()} for x in z.infolist()]
    if s.endswith(('.tgz','.tar.gz','.tar')):
        with tarfile.open(p,'r:*') as t:
            return [{'path':x.name,'bytes':x.size,'type':x.type.decode('ascii','replace') if isinstance(x.type,bytes) else str(x.type)} for x in t]
    raise ValueError('not a supported container')
def bounded_member(path, member, max_bytes=1048576):
    if max_bytes<0: raise ValueError('max_bytes must be nonnegative')
    if '..' in pathlib.PurePosixPath(member).parts: raise ValueError('unsafe traversal path')
    with zipfile.ZipFile(path) as z:
        info=z.getinfo(member)
        with z.open(info) as f: return f.read(max_bytes)
def json_structure(path):
    import ijson
    # Historical full scans contain the authoritative nonfinite counts. This entrypoint
    # refuses bare IEEE constants rather than silently changing their meaning.
    keys=set(); events=collections.Counter()
    with open(path,'rb') as f:
        try:
            for prefix,event,value in ijson.parse(f):
                events[event]+=1
                if event=='map_key': keys.add(str(value))
        except Exception as e:
            return {'status':'NONSTANDARD_JSON_REQUIRES_DECLARED_NORMALIZER','error':type(e).__name__+': '+str(e),'field_names_before_error':sorted(keys),'event_counts_before_error':dict(events)}
    return {'status':'VALID_STANDARD_JSON','field_names':sorted(keys),'event_counts':dict(events)}
def json_tagged_events(path,limit=1000):
    with open(path,'rb') as raw:return json_tagged_stream(raw,limit)
def json_tagged_stream(raw,limit=1000):
    import ijson
    if limit<0:raise ValueError('limit must be nonnegative')
    out=[]; pending=None; truncated=False
    view=NonfiniteTaggedView(raw)
    native=NonfiniteTaggedView.NATIVE.decode()
    def cp(p):return '.'.join(q[len(native):] if q.startswith(native) else q for q in p.split('.'))
    for prefix,event,value in ijson.parse(view):
            if event=='map_key' and value=='__reader_nonfinite__':pending=prefix;continue
            if pending is not None and event=='string' and prefix==pending+'.__reader_nonfinite__':
                if len(out)>=limit:truncated=True;break
                out.append({'path':cp(pending),'type':'nonfinite','value':value});pending=None;continue
            if event in ('string','number','boolean','null'):
                if len(out)>=limit:truncated=True;break
                # Values are returned only to the local caller; source bytes are untouched.
                if event=='string' and isinstance(value,str) and value.startswith(native):value=value[len(native):]
                out.append({'path':cp(prefix),'type':event,'value':value})
    return {'events':out,'nonfinite_counts_seen':dict(view.counts),'complete':not truncated and view.eof,'truncated_by_limit':truncated,'policy':'dedicated nonfinite event type preserves NaN/+Infinity/-Infinity and native strings'}
def zip_json_events(path,member,limit):
    with zipfile.ZipFile(path) as z:
        with z.open(member) as raw:return json_tagged_stream(raw,limit)
def zip_xlsx_structure(path,member):
    with zipfile.ZipFile(path) as z:return xlsx_structure_stream(io.BytesIO(z.read(member)))
def xlsx_structure(path):
    return xlsx_structure_stream(path)
def xlsx_structure_stream(path):
    import openpyxl
    raw=path.read() if hasattr(path,'read') else pathlib.Path(path).read_bytes()
    wb=openpyxl.load_workbook(io.BytesIO(raw),read_only=False,data_only=False,keep_links=False);strict_namespace_compatibility=False
    if not wb.worksheets:
        wb.close();src=zipfile.ZipFile(io.BytesIO(raw));converted=io.BytesIO()
        with src,zipfile.ZipFile(converted,'w') as dst:
            for info in src.infolist():
                data=src.read(info.filename)
                if info.filename.endswith(('.xml','.rels')):
                    data=data.replace(b'http://purl.oclc.org/ooxml/officeDocument/relationships',b'http://schemas.openxmlformats.org/officeDocument/2006/relationships').replace(b'http://purl.oclc.org/ooxml/spreadsheetml/main',b'http://schemas.openxmlformats.org/spreadsheetml/2006/main')
                dst.writestr(info,data)
        converted.seek(0);wb=openpyxl.load_workbook(converted,read_only=False,data_only=False,keep_links=False);strict_namespace_compatibility=True
    out=[]
    for ws in wb.worksheets:
        nonempty=[]; formulas=[];unit_cells=[]
        for row in ws.iter_rows():
            vals=[c.value for c in row]
            if any(v is not None for v in vals) and len(nonempty)<20:nonempty.append({'row':row[0].row,'cells':[str(v) if v is not None else None for v in vals]})
            for c in row:
                if c.data_type=='f' or isinstance(c.value,str) and c.value.startswith('='): formulas.append(c.coordinate)
                if isinstance(c.value,str) and re.search(r'(?i)(\bV\b|volt|\bA\b|amp|mA|Ah|Wh|ohm|°C|deg\s*C|second|\bsec\b|minute|hour|temperature|capacity|current|voltage|time)',c.value):unit_cells.append({'cell':c.coordinate,'text':c.value})
        out.append({'sheet':ws.title,'rows':ws.max_row,'columns':ws.max_column,'first_20_nonempty_rows':nonempty,'formula_count':len(formulas),'formula_cells':formulas,'unit_or_quantity_cells':unit_cells[:500],'unit_cells_truncated':len(unit_cells)>500,'formula_execution':False,'strict_namespace_compatibility':strict_namespace_compatibility})
    wb.close(); return out
def mat73_directory(path,max_nodes=1000000):
    import h5py
    out=[]
    with h5py.File(path,'r') as h:
        def visit(n,o):
            if len(out)>=max_nodes: return 'NODE_LIMIT'
            out.append({'path':n,'kind':'dataset' if isinstance(o,h5py.Dataset) else 'group','shape':list(o.shape) if isinstance(o,h5py.Dataset) else None,'dtype':str(o.dtype) if isinstance(o,h5py.Dataset) else None})
        h.visititems(visit)
    return out
def mat73_slice(path,dataset,start=0,count=1000,max_elements=1000000,max_bytes=67108864):
    import h5py, numpy as np
    if start<0 or count<0 or max_elements<0 or max_bytes<0:raise ValueError('slice bounds must be nonnegative')
    with h5py.File(path,'r') as h:
        d=h[dataset]
        if not isinstance(d,h5py.Dataset): raise ValueError('path is not a dataset')
        if count>max_elements:raise ValueError('requested element count exceeds max_elements')
        if d.ndim==0:a=d[()]
        else:
            trailing=1
            for q in d.shape[1:]:trailing*=q
            if trailing*d.dtype.itemsize>max_bytes:raise ValueError('one axis-0 row exceeds max_bytes; choose a more specific dataset')
            rows=max(1,(count+trailing-1)//trailing) if count else 0
            a=d[start:min(d.shape[0],start+rows)]
        if getattr(a,'dtype',None) is not None and a.dtype.kind=='O': raise ValueError('object/reference datasets require an explicit decoder')
        aa=np.asarray(a).reshape(-1)[:count]
        if aa.size>max_elements or aa.nbytes>max_bytes:raise ValueError(f'slice exceeds limit: elements={aa.size}, bytes={aa.nbytes}')
        return {'dataset':dataset,'source_shape':list(d.shape),'dtype':str(d.dtype),'flat_element_start_axis0':start,'requested_elements':count,'returned_shape':list(aa.shape),'returned_elements':int(aa.size),'returned_bytes':int(aa.nbytes),'max_elements':max_elements,'max_bytes':max_bytes,'values':aa.tolist()}
def numeric_text_profile(path):
    rows=0; widths=collections.Counter(); bad=0
    with open(path,'rt',encoding='utf-8',errors='replace') as f:
        for line in f:
            rows+=1; parts=line.rstrip('\r\n').split('\t'); widths[len(parts)]+=1
            for x in parts:
                if x.strip():
                    try: float(x)
                    except ValueError: bad+=1
    return {'lines':rows,'field_width_counts':dict(widths),'nonnumeric_tokens':bad,'units':'UNKNOWN'}
def npz_metadata(path):
    import numpy as np
    with np.load(path,allow_pickle=False) as z:
        return {k:{'shape':list(z[k].shape),'dtype':str(z[k].dtype),'nan_count':int(np.isnan(z[k]).sum()) if np.issubdtype(z[k].dtype,np.number) else None} for k in z.files}
def main():
    ap=argparse.ArgumentParser(); ap.add_argument('operation',choices=['csv-sample','list-members','member-prefix','json-structure','json-events','zip-json-events','xlsx-structure','zip-xlsx-structure','mat73-directory','mat73-slice','numeric-profile','npz-metadata']); ap.add_argument('path'); ap.add_argument('--member'); ap.add_argument('--dataset'); ap.add_argument('--start',type=int,default=0); ap.add_argument('--limit',type=int,default=1000); ap.add_argument('--max-elements',type=int,default=1000000); ap.add_argument('--max-bytes',type=int,default=67108864); a=ap.parse_args()
    if a.operation=='csv-sample': print(json.dumps(list(csv_rows(a.path,a.limit)),ensure_ascii=False))
    elif a.operation=='list-members': print(json.dumps(list_members(a.path),ensure_ascii=False))
    elif a.operation=='member-prefix': print(bounded_member(a.path,a.member,a.limit).hex())
    elif a.operation=='json-structure': print(json.dumps(json_structure(a.path),ensure_ascii=False))
    elif a.operation=='json-events': print(json.dumps(json_tagged_events(a.path,a.limit),ensure_ascii=False,default=str))
    elif a.operation=='zip-json-events': print(json.dumps(zip_json_events(a.path,a.member,a.limit),ensure_ascii=False,default=str))
    elif a.operation=='xlsx-structure': print(json.dumps(xlsx_structure(a.path),ensure_ascii=False))
    elif a.operation=='zip-xlsx-structure': print(json.dumps(zip_xlsx_structure(a.path,a.member),ensure_ascii=False))
    elif a.operation=='mat73-directory': print(json.dumps(mat73_directory(a.path,a.limit),ensure_ascii=False))
    elif a.operation=='mat73-slice': print(json.dumps(mat73_slice(a.path,a.dataset,a.start,a.limit,a.max_elements,a.max_bytes),ensure_ascii=False))
    elif a.operation=='numeric-profile': print(json.dumps(numeric_text_profile(a.path),ensure_ascii=False))
    else: print(json.dumps(npz_metadata(a.path),ensure_ascii=False))
if __name__=='__main__': main()
