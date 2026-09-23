import json,pathlib,zipfile,io,csv,time,collections,re
import openpyxl
ROOT=pathlib.Path(__file__).parent; RUN=ROOT.parent
def load(p):return [json.loads(x) for x in p.read_text(encoding='utf-8-sig').splitlines() if x.strip()]
src={x['file_id']:x for x in load(RUN/'SOURCE_OBJECTS.jsonl') if x['lane']=='datasets'}
out=ROOT/'CONTAINER_MEMBER_SCHEMAS.jsonl';log=ROOT/'CONTAINER_SCHEMA_PROBE.log';done=set()
if out.exists():done={x['container_file_id'] for x in load(out) if x.get('container_complete_marker')}
DATA={'.csv','.json','.xlsx','.xls','.mat','.txt'}|{f'.{i:03d}' for i in range(100)}
def text_header(b):
 s=b.decode('utf-8-sig','replace'); lines=s.splitlines()[:101]
 if not lines:return {'header':[],'sampled_rows':0}
 delim='\t' if lines[0].count('\t')>lines[0].count(',') else ','
 rr=list(csv.reader(lines,delimiter=delim));return {'header':rr[0],'sampled_rows':max(0,len(rr)-1),'sample_widths':dict(collections.Counter(len(x) for x in rr[1:])),'delimiter':repr(delim)}
def jschema(b):
 try:
  x=json.loads(b)
  return {'parse':'COMPLETE','root_type':type(x).__name__,'top_fields':sorted(x) if isinstance(x,dict) else None,'length':len(x) if hasattr(x,'__len__') else None}
 except Exception as e:return {'parse':'FAILED_OR_PREFIX_ONLY','error':type(e).__name__+': '+str(e)[:200]}
def workbook(b):
 wb=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=False,keep_links=False)
 ss=[]
 for ws in wb.worksheets:
  first=next(ws.iter_rows(),());ss.append({'sheet':ws.title,'rows':ws.max_row,'columns':ws.max_column,'first_row_fields':[str(c.value) if c.value is not None else None for c in first],'formula_in_first_row':any(isinstance(c.value,str) and c.value.startswith('=') for c in first)})
 wb.close();return ss
def one(z,info,parent=None):
 ext=pathlib.Path(info.filename).suffix.lower(); base={'member_path':info.filename,'parent_member_path':parent,'member_bytes':info.file_size,'suffix':ext}
 if info.is_dir():return {**base,'role':'directory','read_status':'NOT_APPLICABLE'}
 if '/._' in info.filename or info.filename.startswith('__MACOSX/') or pathlib.PurePosixPath(info.filename).name.startswith('._'):return {**base,'role':'filesystem_metadata_sidecar','read_status':'EXCLUDED_NON_DATA_RESOURCE_FORK'}
 if ext not in DATA and ext!='.zip':return {**base,'role':'non_data_asset_or_code','read_status':'CLASSIFIED_BY_SUFFIX'}
 if ext=='.xls':return {**base,'role':'data','read_status':'UNSUPPORTED_XLS_READER','error':'openpyxl does not read legacy BIFF .xls; no local xlrd dependency'}
 lim=64*1024*1024
 if info.file_size>lim:
  if ext=='.json':
   try:
    with z.open(info) as s:b=s.read(4*1024*1024)
    keys=sorted(set(x.decode('utf-8','replace') for x in re.findall(rb'"([^"\\]{1,100})"\s*:',b)))
    return {**base,'role':'data','read_status':'BOUNDED_STREAM_SCHEMA_READ','format':'json','schema':{'bytes_read':len(b),'field_names':keys,'scope':'first 4 MiB decompressed; values not retained'},'source_values_emitted':False}
   except Exception as e:return {**base,'role':'data','read_status':'SCHEMA_ERROR','error':type(e).__name__+': '+str(e)[:300]}
  return {**base,'role':'data','read_status':'MEMBER_EXCEEDS_64MIB_IN_MEMORY_LIMIT','error':'requires streaming format-specific follow-up'}
 try:b=z.read(info)
 except Exception as e:return {**base,'role':'data','read_status':'READ_ERROR','error':type(e).__name__+': '+str(e)[:200]}
 try:
  if ext in ('.csv','.txt') or (len(ext)==4 and ext[1:].isdigit()): sch=text_header(b[:2*1024*1024]); kind='delimited_text'
  elif ext=='.json':sch=jschema(b);kind='json'
  elif ext=='.xlsx':sch=workbook(b);kind='xlsx'
  elif ext=='.mat':sch={'signature':'HDF5' if b.startswith(b'\x89HDF') else 'MAT_LEVEL5' if b.startswith(b'MATLAB ') else 'UNKNOWN','bytes_checked':min(len(b),128)};kind='mat'
  elif ext=='.zip':
   with zipfile.ZipFile(io.BytesIO(b)) as n: sch={'nested_members':[one(n,q,info.filename) for q in n.infolist()]}
   kind='nested_zip'
  return {**base,'role':'data_container' if ext=='.zip' else 'data','read_status':'SCHEMA_READ','format':kind,'schema':sch,'source_values_emitted':False}
 except Exception as e:return {**base,'role':'data','read_status':'SCHEMA_ERROR','error':type(e).__name__+': '+str(e)[:300]}
with out.open('a',encoding='utf-8',newline='\n') as fo,log.open('a',encoding='utf-8',buffering=1) as lg:
 for fid,o in src.items():
  if fid in done or o.get('representation')!='ZIP_ARCHIVE':continue
  t=time.time();counts=collections.Counter()
  try:
   with zipfile.ZipFile(o['input_path']) as z:
    for inf in z.infolist():
     r={'container_file_id':fid,**one(z,inf)};counts[r['read_status']]+=1;fo.write(json.dumps(r,ensure_ascii=False,sort_keys=True)+'\n')
   fo.write(json.dumps({'container_file_id':fid,'container_complete_marker':True,'status_counts':dict(counts)},sort_keys=True)+'\n');fo.flush();lg.write(f'{fid}\t{dict(counts)}\t{time.time()-t:.3f}s\n')
  except Exception as e:lg.write(f'{fid}\tCONTAINER_ERROR\t{type(e).__name__}: {e}\n')
