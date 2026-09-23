import json,zipfile,io,collections
from pathlib import Path
from openpyxl import load_workbook
B=Path(r'E:/desn 2000/bsc/data_preparation/CR-DATA-READY-001/20260920T012405_AEST');D=B/'datasets';O=B/'audits'/'doc_reviews_data'
def jl(p):return [json.loads(x) for x in p.read_text(encoding='utf8').splitlines() if x.strip()]
src={x['file_id']:x for x in jl(B/'SOURCE_OBJECTS.jsonl')}; xs=[x for x in jl(D/'CONTAINER_WORKBOOK_CONTRACTS.jsonl') if x['contract_status']=='HEADER_SELECTED']
fams={}
for x in xs:fams.setdefault(tuple(map(str,x.get('selected_fields') or [])),x)
out=[]
for fields,x in fams.items():
 with zipfile.ZipFile(src[x['container_file_id']]['input_path']) as z:data=z.read(x['member_path'])
 wb=load_workbook(io.BytesIO(data),read_only=True,data_only=False);ws=wb[x['sheet']]
 rows=x.get('selected_header_rows') or []; row=rows[0] if len(rows)==1 else None; exact=None; actual=None
 if row:
  actual=[ws.cell(row,c).value for c in range(1,ws.max_column+1)]; exact=[str(v) for v in actual]==list(fields)
 out.append({'fields_family':list(fields),'count':sum(tuple(map(str,q.get('selected_fields') or []))==fields for q in xs),'sample':{'file_id':x['container_file_id'],'member_path':x['member_path'],'sheet':x['sheet'],'selected_header_rows':rows},'actual_selected_row':actual,'exact_row_match':exact,'manual_composite':row is None})
(O/'WORKBOOK_FAMILY_RECHECK.json').write_text(json.dumps({'family_count':len(out),'families':out},ensure_ascii=False,indent=2)+'\n',encoding='utf8')
print(json.dumps({'families':len(out),'exact_rows':sum(q['exact_row_match'] is True for q in out),'mismatch':sum(q['exact_row_match'] is False for q in out),'manual_composite':sum(q['manual_composite'] for q in out)}))

