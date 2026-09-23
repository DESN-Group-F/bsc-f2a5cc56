import json,pathlib,zipfile,io,time,re
import openpyxl
ROOT=pathlib.Path(__file__).parent;RUN=ROOT.parent
def load(p):return [json.loads(x) for x in p.read_text(encoding='utf-8-sig').splitlines() if x.strip()]
src={x['file_id']:x for x in load(RUN/'SOURCE_OBJECTS.jsonl') if x['lane']=='datasets'}
targets=[x for x in load(ROOT/'CONTAINER_MEMBER_SCHEMAS.jsonl') if x.get('read_status')=='SCHEMA_READ' and x.get('suffix')=='.xlsx']
out=ROOT/'CONTAINER_WORKBOOK_PROFILES.jsonl';log=ROOT/'CONTAINER_WORKBOOK_PROFILES.log';done={(x['container_file_id'],x['member_path']) for x in load(out)} if out.exists() else set()
unit_re=re.compile(r'(?i)(\bV\b|volt|\bA\b|amp|mA|Ah|Wh|ohm|°C|deg\s*C|second|\bsec\b|minute|hour|temperature|capacity|current|voltage|time)')
def load_compatible_workbook(b):
 wb=openpyxl.load_workbook(io.BytesIO(b),read_only=False,data_only=False,keep_links=False)
 if wb.worksheets:return wb
 wb.close();src=zipfile.ZipFile(io.BytesIO(b));dst_io=io.BytesIO()
 replacements=((b'http://purl.oclc.org/ooxml/officeDocument/relationships',b'http://schemas.openxmlformats.org/officeDocument/2006/relationships'),(b'http://purl.oclc.org/ooxml/spreadsheetml/main',b'http://schemas.openxmlformats.org/spreadsheetml/2006/main'))
 with src,zipfile.ZipFile(dst_io,'w') as dst:
  for info in src.infolist():
   data=src.read(info.filename)
   if info.filename.endswith(('.xml','.rels')):
    for old,new in replacements:data=data.replace(old,new)
   dst.writestr(info,data)
 dst_io.seek(0);return openpyxl.load_workbook(dst_io,read_only=False,data_only=False,keep_links=False)
with out.open('a',encoding='utf-8',newline='\n') as fo,log.open('a',encoding='utf-8',buffering=1) as lg:
 for n,g in enumerate(targets,1):
  key=(g['container_file_id'],g['member_path'])
  if key in done:continue
  t=time.time();sheets=[];err=None
  try:
   with zipfile.ZipFile(src[key[0]]['input_path']) as z:b=z.read(key[1])
   wb=load_compatible_workbook(b)
   for ws in wb.worksheets:
    nonempty=[];formulas=[];unit_cells=[]
    for row in ws.iter_rows():
     vals=[c.value for c in row]
     if any(v is not None for v in vals) and len(nonempty)<20:nonempty.append({'row':row[0].row,'cells':[str(v) if v is not None else None for v in vals]})
     for c in row:
      if c.data_type=='f' or isinstance(c.value,str) and c.value.startswith('='):formulas.append(c.coordinate)
      if isinstance(c.value,str) and unit_re.search(c.value):unit_cells.append({'cell':c.coordinate,'text':c.value})
    title=' '.join(str(x) for r in nonempty[:5] for x in r['cells'] if x)
    role='documentation_or_disclaimer' if re.search(r'(?i)(disclaimer|information|readme|instruction)',ws.title+' '+title) else 'data_or_analysis_sheet'
    sheets.append({'sheet':ws.title,'role':role,'rows':ws.max_row,'columns':ws.max_column,'first_20_nonempty_rows':nonempty,'merged_ranges':[str(x) for x in ws.merged_cells.ranges],'formula_count':len(formulas),'formula_cells':formulas,'unit_or_quantity_cells':unit_cells[:500],'unit_cells_truncated':len(unit_cells)>500,'formula_execution':False})
   wb.close();status='COMPLETE_WORKBOOK_STRUCTURE'
  except Exception as e:err=type(e).__name__+': '+str(e);status='ERROR'
  fo.write(json.dumps({'container_file_id':key[0],'member_path':key[1],'status':status,'error':err,'sheets':sheets,'elapsed_seconds':round(time.time()-t,3),'source_values_scope':'first 20 nonempty rows and unit/quantity label cells only'},ensure_ascii=False,sort_keys=True)+'\n');fo.flush();lg.write(f'{n}/{len(targets)}\t{key[0]}\t{key[1]}\t{status}\t{time.time()-t:.3f}s\n')
