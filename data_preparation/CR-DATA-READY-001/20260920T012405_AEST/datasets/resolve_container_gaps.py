import json,pathlib,zipfile,tempfile,shutil,time,io
import xlrd
from matio import whosmat
ROOT=pathlib.Path(__file__).parent;RUN=ROOT.parent
def load(p):return [json.loads(x) for x in p.read_text(encoding='utf-8-sig').splitlines() if x.strip()]
src={x['file_id']:x for x in load(RUN/'SOURCE_OBJECTS.jsonl') if x['lane']=='datasets'}
gaps=[x for x in load(ROOT/'CONTAINER_MEMBER_SCHEMAS.jsonl') if x.get('read_status') in ('MEMBER_EXCEEDS_64MIB_IN_MEMORY_LIMIT','UNSUPPORTED_XLS_READER')]
out=[]
for g in gaps:
 fid=g['container_file_id']; member=g['member_path']; t=time.time()
 try:
  with zipfile.ZipFile(src[fid]['input_path']) as outer:
   info=outer.getinfo(member)
   if g['read_status']=='UNSUPPORTED_XLS_READER':
    b=outer.read(info);sheets=[]
    if b.startswith(b'PK'):
     import openpyxl
     book=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=False,keep_links=False)
     for s in book.worksheets:
      first=[];formulas=[]
      for ri,row in enumerate(s.iter_rows(),1):
       if ri==1:first=[str(c.value) if c.value is not None else None for c in row]
       formulas += [c.coordinate for c in row if c.data_type=='f']
      sheets.append({'sheet':s.title,'rows':s.max_row,'columns':s.max_column,'first_row_fields':first,'formula_count':len(formulas),'formula_cells':formulas})
     book.close();decoder='openpyxl_mislabeled_xlsx'
    else:
     book=xlrd.open_workbook(file_contents=b,on_demand=True)
     for s in book.sheets():sheets.append({'sheet':s.name,'rows':s.nrows,'columns':s.ncols,'first_row_fields':[str(s.cell_value(0,c)) if s.nrows else None for c in range(s.ncols)]})
     book.release_resources();decoder='xlrd_biff'
    out.append({**g,'resolution_status':'SCHEMA_READ_WORKBOOK','decoder':decoder,'schema':sheets,'elapsed_seconds':round(time.time()-t,3),'source_values_emitted':False})
   else:
    with tempfile.SpooledTemporaryFile(max_size=16*1024*1024) as sp:
     with outer.open(info) as q:shutil.copyfileobj(q,sp,1024*1024)
     sp.seek(0);nested=[]
     with zipfile.ZipFile(sp) as nz:
      for ni in nz.infolist():
       r={'member_path':ni.filename,'bytes':ni.file_size,'directory':ni.is_dir(),'suffix':pathlib.Path(ni.filename).suffix.lower(),'read_status':'DIRECTORY' if ni.is_dir() else 'SIGNATURE_READ'}
       if not ni.is_dir():
        with nz.open(ni) as q:b=q.read(128)
        r['signature']='MAT_LEVEL5' if b.startswith(b'MATLAB ') else 'HDF5' if b.startswith(b'\x89HDF') else 'TEXT_LIKE' if b and b'\0' not in b else 'BINARY_OR_UNKNOWN'
        if r['suffix']=='.mat' and ni.file_size<=64*1024*1024:
         tmp=tempfile.NamedTemporaryFile(suffix='.mat',delete=False);tmpname=tmp.name
         try:
          with nz.open(ni) as q:shutil.copyfileobj(q,tmp,1024*1024)
          tmp.close();r['variables']={k:v for k,v in whosmat(tmpname).items()};r['read_status']='MAT_VARIABLE_DIRECTORY_READ'
         finally:
          try:pathlib.Path(tmpname).unlink()
          except:pass
       nested.append(r)
    out.append({**g,'resolution_status':'NESTED_ZIP_FULL_DIRECTORY_AND_MAT_SCHEMA','nested_member_count':len(nested),'nested_members':nested,'elapsed_seconds':round(time.time()-t,3),'source_values_emitted':False})
 except Exception as e:out.append({**g,'resolution_status':'RESOLUTION_ERROR','error':type(e).__name__+': '+str(e),'elapsed_seconds':round(time.time()-t,3)})
with (ROOT/'CONTAINER_GAP_RESOLUTIONS.jsonl').open('w',encoding='utf-8',newline='\n') as f:
 for x in out:f.write(json.dumps(x,ensure_ascii=False,sort_keys=True)+'\n')
print(json.dumps({'attempted':len(out),'resolved':sum(x['resolution_status']!='RESOLUTION_ERROR' for x in out),'errors':sum(x['resolution_status']=='RESOLUTION_ERROR' for x in out)}))
